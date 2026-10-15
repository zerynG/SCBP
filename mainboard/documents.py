"""
Генерация документов проекта по стандартам РФ.

Два типа документов с РАЗНЫМИ данными:
  • КП  — Коммерческое предложение: отпускная цена для заказчика
          (себестоимость + наценка + НДС). Реквизиты исполнителя/заказчика,
          таблица работ/услуг, итог без НДС → НДС → итог с НДС, срок, подпись.
  • НМА — Расчёт первоначальной стоимости нематериального актива (ПО),
          ФСБУ 14/2022 + ФСБУ 26/2020: калькуляция фактических затрат
          (ФОТ + страховые взносы + услуги подрядчиков + оборудование),
          т.е. СЕБЕСТОИМОСТЬ создания, без наценки и НДС.

Форматы: PDF (reportlab), Excel (openpyxl), Word (python-docx).
Данные приходят из views._build_project_data (единый калькулятор).
"""
import io
import os
from datetime import datetime
from typing import Any, Dict, List, Optional

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Image as ReportLabImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )
except ImportError:  # pragma: no cover
    pass

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    pass

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # pragma: no cover
    pass


# ---------------------------------------------------------------------------
# Форматирование
# ---------------------------------------------------------------------------

def money(v: Any) -> str:
    """1234567.5 → '1 234 567,50'."""
    try:
        s = f"{float(v):,.2f}"
    except (TypeError, ValueError):
        s = "0.00"
    return s.replace(",", " ").replace(".", ",")


# Палитра
ACCENT = "#6D4AC0"
DARK = "#2C3E50"
GREY = "#7F8C8D"
LIGHT = "#F2F0FA"


class DocumentGenerator:
    """Генератор КП и НМА в PDF/Excel/Word."""

    def __init__(self, project_data: Dict[str, Any],
                 company_settings: Optional[Dict[str, Any]] = None,
                 font_path: Optional[str] = None):
        self.d = project_data
        self.company = company_settings or {}
        self.font_path = font_path
        self.today = datetime.now().strftime('%d.%m.%Y')

    # =====================================================================
    # PDF — общие части
    # =====================================================================

    def _register_fonts(self):
        if self.font_path and os.path.isdir(self.font_path):
            try:
                reg = os.path.join(self.font_path, 'DejaVuSans.ttf')
                bold = os.path.join(self.font_path, 'DejaVuSans-Bold.ttf')
                if os.path.exists(reg):
                    pdfmetrics.registerFont(TTFont('DejaVuSans', reg))
                if os.path.exists(bold):
                    pdfmetrics.registerFont(TTFont('DejaVuSans-Bold', bold))
                    pdfmetrics.registerFontFamily(
                        'DejaVuSans', normal='DejaVuSans', bold='DejaVuSans-Bold',
                        italic='DejaVuSans', boldItalic='DejaVuSans-Bold')
            except Exception:
                pass

    def _fonts(self):
        reg = set(pdfmetrics.getRegisteredFontNames())
        n = 'DejaVuSans' if 'DejaVuSans' in reg else 'Helvetica'
        b = 'DejaVuSans-Bold' if 'DejaVuSans-Bold' in reg else 'Helvetica-Bold'
        return n, b

    def _styles(self):
        n, b = self._fonts()
        s = getSampleStyleSheet()
        s.add(ParagraphStyle('DocTitle', fontName=b, fontSize=16, spaceAfter=4,
                             alignment=1, textColor=colors.HexColor(DARK)))
        s.add(ParagraphStyle('DocSubtitle', fontName=n, fontSize=10, spaceAfter=14,
                             alignment=1, textColor=colors.HexColor(GREY)))
        s.add(ParagraphStyle('H2', fontName=b, fontSize=12, spaceBefore=10, spaceAfter=6,
                             textColor=colors.HexColor(ACCENT)))
        s.add(ParagraphStyle('Body', fontName=n, fontSize=9.5, spaceAfter=4,
                             textColor=colors.HexColor(DARK), leading=14))
        s.add(ParagraphStyle('Small', fontName=n, fontSize=8, textColor=colors.HexColor(GREY)))
        s.add(ParagraphStyle('CellL', fontName=n, fontSize=8.5, textColor=colors.HexColor(DARK), leading=11))
        s.add(ParagraphStyle('CellR', fontName=n, fontSize=8.5, alignment=2, textColor=colors.HexColor(DARK)))
        s.add(ParagraphStyle('CellHdr', fontName=b, fontSize=8.5, alignment=1, textColor=colors.white))
        return s

    def _company_block(self, st):
        c = self.company
        lines = [f"<b>{c.get('name', '')}</b>"]
        rek = []
        if c.get('inn'):
            rek.append(f"ИНН {c['inn']}")
        if c.get('kpp'):
            rek.append(f"КПП {c['kpp']}")
        if c.get('ogrn'):
            rek.append(f"ОГРН {c['ogrn']}")
        if rek:
            lines.append(', '.join(rek))
        if c.get('address'):
            lines.append(c['address'])
        contacts = [x for x in (c.get('phone'), c.get('email')) if x]
        if contacts:
            lines.append(', '.join(contacts))
        return Paragraph('<br/>'.join(lines), st['Small'])

    def _signature_block(self, st):
        c = self.company
        pos = c.get('ceo_position') or 'Руководитель'
        who = c.get('ceo_full_name') or '______________________'
        html = (f"{pos} ____________________ / {who} /<br/>"
                f"<font color='{GREY}'>М.П.</font>")
        return Paragraph(html, st['Body'])

    def _pdf(self, build_fn) -> io.BytesIO:
        self._register_fonts()
        st = self._styles()
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm,
                                leftMargin=18 * mm, rightMargin=18 * mm)
        elements: List[Any] = []
        # Логотип (если задан)
        logo = self.company.get('logo_path')
        if logo and os.path.exists(logo):
            try:
                img = ReportLabImage(logo, width=45 * mm, height=15 * mm)
                img.hAlign = 'LEFT'
                elements.append(img)
                elements.append(Spacer(1, 6))
            except Exception:
                pass
        elements.append(self._company_block(st))
        elements.append(Spacer(1, 10))
        build_fn(elements, st)
        elements.append(Spacer(1, 16))
        elements.append(self._signature_block(st))
        elements.append(Spacer(1, 8))
        elements.append(Paragraph(f"Документ сформирован {self.today}", st['Small']))
        doc.build(elements)
        buf.seek(0)
        return buf

    def _table_style(self, ncols, money_cols=()):
        n, b = self._fonts()
        ts = [
            ('FONTNAME', (0, 0), (-1, -1), n),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(ACCENT)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), b),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBC6DD')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor(LIGHT)]),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ]
        for c in money_cols:
            ts.append(('ALIGN', (c, 1), (c, -1), 'RIGHT'))
        return TableStyle(ts)

    # =====================================================================
    # КП — Коммерческое предложение (PDF)
    # =====================================================================

    def _build_kp_pdf(self, el, st):
        d = self.d
        el.append(Paragraph(f"Коммерческое предложение № {d.get('doc_number', '')} от {self.today}", st['DocTitle']))
        el.append(Paragraph(f"по проекту «{d.get('name', '')}»", st['DocSubtitle']))
        el.append(Paragraph(f"<b>Заказчик:</b> {d.get('customer', {}).get('name', '—')}", st['Body']))
        period = ''
        if d.get('start_date') and d.get('start_date') != '—':
            period = f"  &nbsp;&nbsp;<b>Срок:</b> {d.get('start_date')} — {d.get('end_date')}"
        el.append(Paragraph(f"Предлагаем выполнение работ/услуг по проекту.{period}", st['Body']))
        el.append(Spacer(1, 8))

        # Таблица работ
        header = ['№', 'Наименование работ / услуг', 'Ед.', 'Кол-во', 'Цена, ₽', 'Сумма, ₽']
        rows = [[Paragraph(h, st['CellHdr']) for h in header]]
        for it in d.get('kp_items', []):
            rows.append([
                Paragraph(str(it['num']), st['CellL']),
                Paragraph(it['name'], st['CellL']),
                Paragraph(it['unit'], st['CellL']),
                Paragraph(str(it['qty']), st['CellR']),
                Paragraph(money(it['price']), st['CellR']),
                Paragraph(money(it['sum']), st['CellR']),
            ])
        if len(rows) == 1:
            rows.append([Paragraph('—', st['CellL'])] + [Paragraph('', st['CellL']) for _ in range(5)])
        t = Table(rows, colWidths=[10 * mm, 72 * mm, 13 * mm, 21 * mm, 28 * mm, 30 * mm])
        t.setStyle(self._table_style(6, money_cols=(3, 4, 5)))
        el.append(t)
        el.append(Spacer(1, 8))

        # Итоги
        el.append(self._totals_table_pdf(st, [
            ('Итого без НДС:', money(d.get('final_price', 0))),
            (f"НДС {('%g' % d.get('vat_rate', 0))}%:", money(d.get('vat_amount', 0))),
            ('Итого к оплате (с НДС):', money(d.get('price_with_vat', 0))),
        ], bold_last=True))

    def _totals_table_pdf(self, st, pairs, bold_last=False):
        n, b = self._fonts()
        rows = []
        for label, val in pairs:
            rows.append([Paragraph(label, st['CellR']), Paragraph(val + ' ₽', st['CellR'])])
        t = Table(rows, colWidths=[120 * mm, 54 * mm], hAlign='RIGHT')
        ts = [
            ('FONTNAME', (0, 0), (-1, -1), n),
            ('FONTSIZE', (0, 0), (-1, -1), 9.5),
            ('ALIGN', (0, 0), (-1, -1), 'RIGHT'),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('LINEABOVE', (0, -1), (-1, -1), 0.7, colors.HexColor(ACCENT)),
        ]
        if bold_last:
            ts += [('FONTNAME', (0, -1), (-1, -1), b), ('FONTSIZE', (0, -1), (-1, -1), 11),
                   ('TEXTCOLOR', (0, -1), (-1, -1), colors.HexColor(ACCENT))]
        t.setStyle(TableStyle(ts))
        return t

    def generate_proposal_pdf(self) -> io.BytesIO:
        return self._pdf(self._build_kp_pdf)

    # =====================================================================
    # НМА — расчёт первоначальной стоимости (PDF)
    # =====================================================================

    def _build_nma_pdf(self, el, st):
        d = self.d
        el.append(Paragraph(f"Расчёт первоначальной стоимости НМА № {d.get('doc_number', '')}", st['DocTitle']))
        el.append(Paragraph("нематериальный актив (ПО) · ФСБУ 14/2022, ФСБУ 26/2020", st['DocSubtitle']))
        el.append(Paragraph(f"<b>Объект (проект):</b> {d.get('name', '')}", st['Body']))
        el.append(Paragraph(f"<b>Заказчик:</b> {d.get('customer', {}).get('name', '—')}", st['Body']))
        el.append(Spacer(1, 6))

        # Калькуляция по статьям затрат
        el.append(Paragraph("Калькуляция фактических затрат (стадия разработки)", st['H2']))
        header = ['№', 'Статья затрат', 'Сумма, ₽']
        rows = [[Paragraph(h, st['CellHdr']) for h in header]]
        for i, b in enumerate(d.get('nma_breakdown', []), 1):
            rows.append([
                Paragraph(str(i), st['CellL']),
                Paragraph(b['article'], st['CellL']),
                Paragraph(money(b['amount']), st['CellR']),
            ])
        rows.append([
            Paragraph('', st['CellL']),
            Paragraph('<b>Первоначальная стоимость НМА (итого):</b>', st['CellL']),
            Paragraph('<b>' + money(d.get('nma_total', 0)) + '</b>', st['CellR']),
        ])
        t = Table(rows, colWidths=[10 * mm, 134 * mm, 30 * mm])
        ts = self._table_style(3, money_cols=(2,))
        ts.add('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(LIGHT))
        ts.add('LINEABOVE', (0, -1), (-1, -1), 0.7, colors.HexColor(ACCENT))
        t.setStyle(ts)
        el.append(t)
        el.append(Spacer(1, 6))
        el.append(Paragraph(
            "Стоимость сформирована по фактическим затратам на стадии разработки "
            "(оплата труда и страховые взносы, услуги соисполнителей, оборудование). "
            "НДС и наценка в первоначальную стоимость НМА не включаются. "
            "Принятие к учёту: Дт 04 «НМА» Кт 08-5 «Вложения во внеоборотные активы».",
            st['Small']))

        # Детализация по ресурсам
        el.append(Spacer(1, 10))
        el.append(Paragraph("Детализация по ресурсам", st['H2']))
        header2 = ['Ресурс', 'Тип', 'Услуга', 'Себест., ₽']
        rows2 = [[Paragraph(h, st['CellHdr']) for h in header2]]
        for r in d.get('resources', []):
            rows2.append([
                Paragraph(r['name'], st['CellL']),
                Paragraph(r['type'], st['CellL']),
                Paragraph(r['service'] or '—', st['CellL']),
                Paragraph(money(r['cost']), st['CellR']),
            ])
        t2 = Table(rows2, colWidths=[56 * mm, 30 * mm, 56 * mm, 32 * mm])
        t2.setStyle(self._table_style(4, money_cols=(3,)))
        el.append(t2)

    def generate_nma_pdf(self) -> io.BytesIO:
        return self._pdf(self._build_nma_pdf)

    # =====================================================================
    # Excel — общие части
    # =====================================================================

    def _xl_styles(self):
        return {
            'title': Font(name='Calibri', size=16, bold=True, color='2C3E50'),
            'sub': Font(name='Calibri', size=10, color='7F8C8D'),
            'h2': Font(name='Calibri', size=12, bold=True, color='6D4AC0'),
            'hdr': Font(name='Calibri', size=10, bold=True, color='FFFFFF'),
            'bold': Font(name='Calibri', size=10, bold=True, color='2C3E50'),
            'norm': Font(name='Calibri', size=10, color='2C3E50'),
            'small': Font(name='Calibri', size=8, color='7F8C8D'),
            'hdr_fill': PatternFill('solid', fgColor='6D4AC0'),
            'tot_fill': PatternFill('solid', fgColor='F2F0FA'),
            'center': Alignment(horizontal='center', vertical='center', wrap_text=True),
            'left': Alignment(horizontal='left', vertical='center', wrap_text=True),
            'right': Alignment(horizontal='right', vertical='center'),
            'border': Border(*[Side(style='thin', color='CBC6DD')] * 4),
        }

    def _xl_company(self, ws, s, row):
        c = self.company
        ws.cell(row=row, column=1, value=c.get('name', '')).font = s['bold']
        row += 1
        rek = []
        for k, lbl in (('inn', 'ИНН'), ('kpp', 'КПП'), ('ogrn', 'ОГРН')):
            if c.get(k):
                rek.append(f"{lbl} {c[k]}")
        line = ', '.join(rek)
        for extra in (c.get('address'), ', '.join(x for x in (c.get('phone'), c.get('email')) if x)):
            if extra:
                line = (line + ('; ' if line else '') + extra)
        if line:
            ws.cell(row=row, column=1, value=line).font = s['small']
            row += 1
        return row + 1

    def _xl_table(self, ws, s, start_row, header, data, widths, money_cols, total_row=None):
        col_n = len(header)
        for j, h in enumerate(header, 1):
            cell = ws.cell(row=start_row, column=j, value=h)
            cell.font = s['hdr']; cell.fill = s['hdr_fill']
            cell.alignment = s['center']; cell.border = s['border']
        r = start_row + 1
        for rowvals in data:
            for j, v in enumerate(rowvals, 1):
                cell = ws.cell(row=r, column=j, value=v)
                cell.font = s['norm']; cell.border = s['border']
                if (j - 1) in money_cols:
                    cell.alignment = s['right']; cell.number_format = '# ##0.00'
                else:
                    cell.alignment = s['left']
            r += 1
        if total_row is not None:
            for j, v in enumerate(total_row, 1):
                cell = ws.cell(row=r, column=j, value=v)
                cell.font = s['bold']; cell.fill = s['tot_fill']; cell.border = s['border']
                if (j - 1) in money_cols:
                    cell.alignment = s['right']; cell.number_format = '# ##0.00'
                else:
                    cell.alignment = s['left']
            r += 1
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        return r

    def _xl_totals(self, ws, s, row, pairs, ncols):
        lbl_col = max(1, ncols - 1)
        for label, val in pairs:
            ws.cell(row=row, column=lbl_col, value=label).font = s['bold']
            ws.cell(row=row, column=lbl_col).alignment = s['right']
            vc = ws.cell(row=row, column=ncols, value=val)
            vc.font = s['bold']; vc.alignment = s['right']; vc.number_format = '# ##0.00'
            row += 1
        return row

    def generate_proposal_excel(self) -> io.BytesIO:
        s = self._xl_styles()
        wb = Workbook(); ws = wb.active; ws.title = 'КП'
        d = self.d
        ws.merge_cells('A1:F1')
        ws['A1'] = f"Коммерческое предложение № {d.get('doc_number','')} от {self.today}"
        ws['A1'].font = s['title']; ws['A1'].alignment = s['center']
        row = self._xl_company(ws, s, 3)
        ws.cell(row=row, column=1, value=f"Проект: {d.get('name','')}").font = s['norm']; row += 1
        ws.cell(row=row, column=1, value=f"Заказчик: {d.get('customer',{}).get('name','—')}").font = s['norm']; row += 2

        header = ['№', 'Наименование работ / услуг', 'Ед.', 'Кол-во', 'Цена', 'Сумма']
        data = [[it['num'], it['name'], it['unit'], it['qty'], it['price'], it['sum']]
                for it in d.get('kp_items', [])]
        row = self._xl_table(ws, s, row, header, data, [6, 50, 8, 10, 16, 18], money_cols={4, 5})
        row += 1
        self._xl_totals(ws, s, row, [
            ('Итого без НДС:', d.get('final_price', 0)),
            (f"НДС {('%g' % d.get('vat_rate',0))}%:", d.get('vat_amount', 0)),
            ('Итого с НДС:', d.get('price_with_vat', 0)),
        ], ncols=6)
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        return buf

    def generate_nma_excel(self) -> io.BytesIO:
        s = self._xl_styles()
        wb = Workbook(); ws = wb.active; ws.title = 'НМА'
        d = self.d
        ws.merge_cells('A1:C1')
        ws['A1'] = f"Расчёт первоначальной стоимости НМА № {d.get('doc_number','')}"
        ws['A1'].font = s['title']; ws['A1'].alignment = s['center']
        ws.merge_cells('A2:C2')
        ws['A2'] = "нематериальный актив (ПО) · ФСБУ 14/2022, ФСБУ 26/2020"
        ws['A2'].font = s['sub']; ws['A2'].alignment = s['center']
        row = self._xl_company(ws, s, 4)
        ws.cell(row=row, column=1, value=f"Объект (проект): {d.get('name','')}").font = s['norm']; row += 1
        ws.cell(row=row, column=1, value=f"Заказчик: {d.get('customer',{}).get('name','—')}").font = s['norm']; row += 2

        header = ['№', 'Статья затрат', 'Сумма']
        data = [[i, b['article'], b['amount']] for i, b in enumerate(d.get('nma_breakdown', []), 1)]
        total = ['', 'Первоначальная стоимость НМА (итого)', d.get('nma_total', 0)]
        row = self._xl_table(ws, s, row, header, data, [6, 60, 18], money_cols={2}, total_row=total)
        row += 1
        ws.cell(row=row, column=1, value=(
            "НДС и наценка не включаются. Принятие к учёту: Дт 04 Кт 08-5.")).font = s['small']
        row += 2
        # Детализация
        ws.cell(row=row, column=1, value='Детализация по ресурсам').font = s['h2']; row += 1
        header2 = ['Ресурс', 'Тип', 'Услуга', 'Себестоимость']
        data2 = [[r['name'], r['type'], r['service'] or '—', r['cost']] for r in d.get('resources', [])]
        self._xl_table(ws, s, row, header2, data2, [34, 18, 30, 18], money_cols={3})
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        return buf

    # =====================================================================
    # Word — общие части
    # =====================================================================

    def _docx_base(self):
        doc = Document()
        normal = doc.styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(10)
        c = self.company
        p = doc.add_paragraph(); r = p.add_run(c.get('name', '')); r.bold = True; r.font.size = Pt(11)
        rek = []
        for k, lbl in (('inn', 'ИНН'), ('kpp', 'КПП'), ('ogrn', 'ОГРН')):
            if c.get(k):
                rek.append(f"{lbl} {c[k]}")
        line = ', '.join(rek)
        for extra in (c.get('address'), ', '.join(x for x in (c.get('phone'), c.get('email')) if x)):
            if extra:
                line = line + ('; ' if line else '') + extra
        if line:
            sp = doc.add_paragraph(); sr = sp.add_run(line); sr.font.size = Pt(8); sr.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
        return doc

    def _docx_table(self, doc, header, data, money_cols, total_row=None, widths=None):
        table = doc.add_table(rows=1, cols=len(header)); table.style = 'Light Grid Accent 1'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = table.rows[0].cells
        for j, h in enumerate(header):
            hdr[j].text = h
            for par in hdr[j].paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.bold = True; run.font.size = Pt(9)
        for rowvals in data:
            cells = table.add_row().cells
            for j, v in enumerate(rowvals):
                cells[j].text = money(v) if j in money_cols else str(v)
                for par in cells[j].paragraphs:
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT if j in money_cols else WD_ALIGN_PARAGRAPH.LEFT
                    for run in par.runs:
                        run.font.size = Pt(9)
        if total_row is not None:
            cells = table.add_row().cells
            for j, v in enumerate(total_row):
                cells[j].text = money(v) if j in money_cols else str(v)
                for par in cells[j].paragraphs:
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT if j in money_cols else WD_ALIGN_PARAGRAPH.LEFT
                    for run in par.runs:
                        run.bold = True; run.font.size = Pt(9)
        if widths:
            for row in table.rows:
                for j, w in enumerate(widths):
                    row.cells[j].width = Cm(w)
        return table

    def _docx_totals(self, doc, pairs):
        for label, val in pairs:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(f"{label} {money(val)} ₽"); r.bold = True
        return doc

    def _docx_signature(self, doc):
        c = self.company
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.add_run(f"{c.get('ceo_position') or 'Руководитель'} ____________________ / "
                  f"{c.get('ceo_full_name') or '________________'} /")
        doc.add_paragraph('М.П.').runs[0].font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
        sp = doc.add_paragraph(f"Документ сформирован {self.today}")
        sp.runs[0].font.size = Pt(8); sp.runs[0].font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

    def generate_proposal_word(self) -> io.BytesIO:
        d = self.d
        doc = self._docx_base()
        h = doc.add_heading(f"Коммерческое предложение № {d.get('doc_number','')} от {self.today}", level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub = doc.add_paragraph(f"по проекту «{d.get('name','')}»"); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Заказчик: {d.get('customer',{}).get('name','—')}")
        if d.get('start_date') and d.get('start_date') != '—':
            doc.add_paragraph(f"Срок выполнения: {d.get('start_date')} — {d.get('end_date')}")
        doc.add_paragraph("Предлагаем выполнение следующих работ/услуг:")
        header = ['№', 'Наименование работ / услуг', 'Ед.', 'Кол-во', 'Цена, ₽', 'Сумма, ₽']
        data = [[it['num'], it['name'], it['unit'], it['qty'], it['price'], it['sum']]
                for it in d.get('kp_items', [])]
        self._docx_table(doc, header, data, money_cols={4, 5}, widths=[1, 7.5, 1.3, 1.6, 2.6, 2.8])
        doc.add_paragraph()
        self._docx_totals(doc, [
            ('Итого без НДС:', d.get('final_price', 0)),
            (f"НДС {('%g' % d.get('vat_rate',0))}%:", d.get('vat_amount', 0)),
            ('Итого к оплате (с НДС):', d.get('price_with_vat', 0)),
        ])
        self._docx_signature(doc)
        buf = io.BytesIO(); doc.save(buf); buf.seek(0)
        return buf

    def generate_nma_word(self) -> io.BytesIO:
        d = self.d
        doc = self._docx_base()
        h = doc.add_heading(f"Расчёт первоначальной стоимости НМА № {d.get('doc_number','')}", level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub = doc.add_paragraph("нематериальный актив (ПО) · ФСБУ 14/2022, ФСБУ 26/2020")
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Объект (проект): {d.get('name','')}")
        doc.add_paragraph(f"Заказчик: {d.get('customer',{}).get('name','—')}")
        doc.add_heading('Калькуляция фактических затрат (стадия разработки)', level=2)
        header = ['№', 'Статья затрат', 'Сумма, ₽']
        data = [[i, b['article'], b['amount']] for i, b in enumerate(d.get('nma_breakdown', []), 1)]
        total = ['', 'Первоначальная стоимость НМА (итого)', d.get('nma_total', 0)]
        self._docx_table(doc, header, data, money_cols={2}, total_row=total, widths=[1, 11, 3])
        note = doc.add_paragraph(
            "НДС и наценка в первоначальную стоимость НМА не включаются. "
            "Принятие к учёту: Дт 04 «НМА» Кт 08-5.")
        note.runs[0].font.size = Pt(8); note.runs[0].font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
        doc.add_heading('Детализация по ресурсам', level=2)
        header2 = ['Ресурс', 'Тип', 'Услуга', 'Себестоимость, ₽']
        data2 = [[r['name'], r['type'], r['service'] or '—', r['cost']] for r in d.get('resources', [])]
        self._docx_table(doc, header2, data2, money_cols={3}, widths=[5.5, 3, 4.5, 3])
        self._docx_signature(doc)
        buf = io.BytesIO(); doc.save(buf); buf.seek(0)
        return buf
