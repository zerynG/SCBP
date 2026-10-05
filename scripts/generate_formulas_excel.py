"""
Генератор `formulas.xlsx` — справочника по формулам расчёта SCBP.

Источник истины — `mainboard/calculator.py`. Здесь описываем формулы и
ссылки на функции.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def build_rows():
    return [
        # Календарь
        {
            "id": "CAL-1",
            "name": "Рабочие дни в периоде",
            "description": "Подсчёт рабочих дней по JSON-календарю. Если для года календарь недоступен — fallback на Пн-Пт.",
            "inputs": "start, end, calendars",
            "outputs": "day_count",
            "formula": "day_count = Σ is_work_day(d) for d in [start..end]",
            "source": "mainboard/calculator.py:count_business_days",
        },
        {
            "id": "CAL-2",
            "name": "Рабочих дней в месяце",
            "description": "Количество рабочих дней в месяце ref_date. Если календарь дал 0 — fallback 22.",
            "inputs": "ref_date, calendars, fallback=22",
            "outputs": "days",
            "formula": "days = count_business_days(1st_of_month, last_of_month, calendars) or 22",
            "source": "mainboard/calculator.py:working_days_in_month",
        },
        {
            "id": "CAL-3",
            "name": "Недоступные рабочие дни",
            "description": "Сколько рабочих дней (по JSON-календарю) сотрудник был недоступен (отпуск/другой проект).",
            "inputs": "schedule_entries, period_start, period_end, calendars",
            "outputs": "count",
            "formula": "count = |{d in unavail_set : is_work_day(d, calendars)}|",
            "source": "mainboard/calculator.py:count_unavailable_business_days",
        },
        # Employee
        {
            "id": "EMP-1",
            "name": "Себестоимость сотрудника за период",
            "description": "С учётом календаря и недоступных дней. Месяцы = (раб_дни - недоступные) / раб_дни_в_месяце.",
            "inputs": "monthly_salary, start, end, calendars, unavailable_days",
            "outputs": "base_no_tax, work_days, months",
            "formula": "base_no_tax = monthly_salary × (work_days - unavailable) / working_days_in_month(start)",
            "source": "mainboard/calculator.py:calc_employee_base",
        },
        {
            "id": "EMP-2",
            "name": "Себестоимость сотрудника без дат",
            "description": "Если start/end не заданы — полная месячная зарплата (как fallback).",
            "inputs": "monthly_salary",
            "outputs": "base_no_tax",
            "formula": "base_no_tax = monthly_salary",
            "source": "mainboard/calculator.py:calc_employee_base",
        },
        # Contractor
        {
            "id": "CTR-1",
            "name": "Себестоимость исполнителя (ГПХ/НПД)",
            "description": "Часы передаются явно (включая 0) → rate × часы. Иначе — rate × estimated_hours.",
            "inputs": "hourly_rate, working_hours | estimated_hours | estimated_cost",
            "outputs": "base_no_tax",
            "formula": (
                "if working_hours not None: rate × working_hours; "
                "elif rate > 0: rate × estimated_hours; "
                "else: estimated_cost"
            ),
            "source": "mainboard/calculator.py:calc_contractor_base",
        },
        # Subcontractor
        {
            "id": "SUB-1",
            "name": "Себестоимость подрядчика",
            "description": "Почасовая: rate × часы (working или estimated). Фиксированная: contract_amount.",
            "inputs": "pricing_type, contract_amount | (hourly_rate, working_hours | estimated_hours)",
            "outputs": "base_no_tax",
            "formula": (
                "if pricing_type=='hourly' and rate>0: rate × (working_hours or estimated_hours); "
                "else: contract_amount"
            ),
            "source": "mainboard/calculator.py:calc_subcontractor_base",
        },
        # Equipment
        {
            "id": "EQ-1",
            "name": "Оборудование — месячная ставка",
            "description": "Округлённое вверх количество календарных месяцев. Правило: разница месяцев + 1, если end.day ≥ start.day. min(1).",
            "inputs": "operational_cost, start, end, quantity",
            "outputs": "base_no_tax, months_count",
            "formula": (
                "months = (end.year-start.year)×12 + (end.month-start.month); "
                "if end.day ≥ start.day: months += 1; months = max(1, months); "
                "base_no_tax = operational_cost × months × quantity"
            ),
            "source": "mainboard/calculator.py:calc_equipment_base / equipment_months",
        },
        {
            "id": "EQ-2a",
            "name": "Оборудование — почасовая, база 'calendar' (по умолчанию)",
            "description": "Календарные часы: каждые сутки = 24 часа. Подходит для постоянной аренды (сервера, ЦОД и т.п.).",
            "inputs": "operational_cost, start, end, quantity",
            "outputs": "base_no_tax, calendar_hours",
            "formula": "base_no_tax = operational_cost × ((end-start).days + 1) × 24 × quantity",
            "source": "mainboard/calculator.py:calc_equipment_base (hour_basis='calendar')",
        },
        {
            "id": "EQ-2b",
            "name": "Оборудование — почасовая, база 'working'",
            "description": "Рабочие часы: рабочих_дней (по календарю) × КЧЧ (часов в дне). Включается в настройках проекта.",
            "inputs": "operational_cost, start, end, quantity, calendars, work_hours_per_day",
            "outputs": "base_no_tax, calendar_hours",
            "formula": "base_no_tax = operational_cost × business_days(start,end) × work_hours_per_day × quantity",
            "source": "mainboard/calculator.py:calc_equipment_base (hour_basis='working')",
        },
        {
            "id": "EQ-3",
            "name": "Оборудование — поминутная",
            "description": "days × 24 × 60 минут.",
            "inputs": "operational_cost, start, end, quantity",
            "outputs": "base_no_tax",
            "formula": "base_no_tax = operational_cost × days × 24 × 60 × quantity",
            "source": "mainboard/calculator.py:calc_equipment_base",
        },
        {
            "id": "EQ-4",
            "name": "Оборудование без дат",
            "description": "Fallback к total_cost (cost × quantity) когда start/end не заданы.",
            "inputs": "total_cost",
            "outputs": "base_no_tax",
            "formula": "base_no_tax = total_cost",
            "source": "mainboard/calculator.py:calc_equipment_base",
        },
        # Tax / margin
        {
            "id": "TAX-1",
            "name": "Применение налога и маржи",
            "description": "Единое правило для всех типов ресурсов. Применяется ровно один раз.",
            "inputs": "base_no_tax, project.tax_rate, margin_percent",
            "outputs": "cost, final_cost",
            "formula": (
                "cost = base_no_tax × (1 + project.tax_rate/100); "
                "final_cost = cost × (1 + margin_percent/100)"
            ),
            "source": "mainboard/calculator.py:apply_tax_and_margin",
        },
        {
            "id": "TAX-2",
            "name": "Выбор маржи ресурса vs проекта",
            "description": (
                "override.resource_margin = NULL → берём project.margin. "
                "= 0 → строгий 0% (явно). Иначе — override."
            ),
            "inputs": "override_margin, project_margin",
            "outputs": "margin_percent",
            "formula": "margin = override_margin if override_margin is not None else project_margin",
            "source": "mainboard/calculator.py:resolve_margin",
        },
        # Настройки проекта (новые)
        {
            "id": "CFG-1",
            "name": "Настройка: Часов в рабочем дне (КЧЧ)",
            "description": "Применяется для отображения «Часы» сотрудника (workDays × КЧЧ) и для оборудования с базой 'working'.",
            "inputs": "Project.work_hours_per_day",
            "outputs": "—",
            "formula": "default = 8; диапазон 0.01..24",
            "source": "mainboard/models.py:Project.work_hours_per_day",
        },
        {
            "id": "CFG-2",
            "name": "Настройка: База часов оборудования",
            "description": "Переключатель между календарными (×24) и рабочими (×КЧЧ) часами для pricing_type='hour'.",
            "inputs": "Project.equipment_hour_basis",
            "outputs": "—",
            "formula": "default = 'calendar'; choices = ['calendar', 'working']",
            "source": "mainboard/models.py:Project.equipment_hour_basis",
        },
        {
            "id": "CFG-3",
            "name": "Настройка: Fallback рабочих дней в месяце",
            "description": "Используется, если JSON-календарь недоступен или вернул 0 для месяца.",
            "inputs": "Project.working_days_per_month_fallback",
            "outputs": "—",
            "formula": "default = 22; диапазон 1..31",
            "source": "mainboard/models.py:Project.working_days_per_month_fallback",
        },

        # Project totals
        {
            "id": "PRJ-1",
            "name": "Себестоимость проекта",
            "description": "Сумма cost всех ресурсов (с налогом, без маржи).",
            "inputs": "rows[].cost",
            "outputs": "base_cost",
            "formula": "base_cost = Σ row.cost",
            "source": "mainboard/calculator.py:calc_project / ProjectTotals",
        },
        {
            "id": "PRJ-2",
            "name": "Итоговая стоимость проекта",
            "description": "Сумма final_cost всех ресурсов. Налог и маржа уже применены на уровне строки — повторно не дублируются.",
            "inputs": "rows[].final_cost",
            "outputs": "final_price",
            "formula": "final_price = Σ row.final_cost",
            "source": "mainboard/calculator.py:calc_project",
        },
        {
            "id": "PRJ-3",
            "name": "Чистая прибыль",
            "description": "Разница итога и себестоимости.",
            "inputs": "final_price, base_cost",
            "outputs": "net_profit",
            "formula": "net_profit = final_price - base_cost",
            "source": "mainboard/calculator.py:calc_project",
        },
    ]


def write_excel(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Формулы"

    headers = ["ID", "Название", "Описание", "Входы", "Выходы", "Формула", "Источник"]
    header_fill = PatternFill(start_color="8B5CF6", end_color="8B5CF6", fill_type="solid")
    header_font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    cell_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    for idx, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center
        c.border = thin

    rows = build_rows()
    for r_idx, item in enumerate(rows, start=2):
        ws.cell(row=r_idx, column=1, value=item["id"]).font = cell_font
        ws.cell(row=r_idx, column=2, value=item["name"]).font = cell_font
        ws.cell(row=r_idx, column=3, value=item["description"]).font = cell_font
        ws.cell(row=r_idx, column=4, value=item["inputs"]).font = cell_font
        ws.cell(row=r_idx, column=5, value=item["outputs"]).font = cell_font
        ws.cell(row=r_idx, column=6, value=item["formula"]).font = cell_font
        ws.cell(row=r_idx, column=7, value=item["source"]).font = cell_font
        for col in range(1, 8):
            cell = ws.cell(row=r_idx, column=col)
            cell.alignment = center if col == 1 else left
            cell.border = thin

    widths = {"A": 12, "B": 36, "C": 56, "D": 36, "E": 22, "F": 70, "G": 50}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    wb.save(path)


if __name__ == "__main__":
    write_excel("formulas.xlsx")
